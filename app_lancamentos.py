import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows
from datetime import datetime
import os
import warnings
from openpyxl import Workbook

# Ignorar avisos do openpyxl
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# Configuração da página
st.set_page_config(
    page_title="Registro de Lançamentos Financeiros",
    page_icon="💰",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Caminho do arquivo Excel
ARQUIVO_EXCEL = "planilha_financeira.xlsx"

# Função para carregar categorias e métodos da planilha
@st.cache_data
def carregar_configuracoes():
    """Carrega as configurações da aba Configurações"""
    try:
        df_config = pd.read_excel(ARQUIVO_EXCEL, sheet_name='Configurações', header=None)
        
        # Extrair categorias
        necessidades = df_config.iloc[1:, 0].dropna().tolist()
        desejos = df_config.iloc[1:, 1].dropna().tolist()
        investimentos = df_config.iloc[1:, 2].dropna().tolist()
        
        # Extrair métodos
        metodos = df_config.iloc[1:, 3].dropna().tolist()
        
        # Extrair tipos
        tipos = df_config.iloc[1:, 4].dropna().tolist()
        
        todas_categorias = necessidades + desejos + investimentos
        
        return {
            'necessidades': necessidades,
            'desejos': desejos,
            'investimentos': investimentos,
            'todas_categorias': todas_categorias,
            'metodos': metodos,
            'tipos': tipos
        }
    except Exception as e:
        st.error(f"Erro ao carregar configurações: {e}")
        return None

# Função para carregar lançamentos existentes
def carregar_lancamentos():
    """Carrega os lançamentos da aba Lançamentos"""
    try:
        df = pd.read_excel(ARQUIVO_EXCEL, sheet_name='Lançamentos')
        
        # Se a coluna Contas não existe, criar com valor padrão "Nubank"
        if 'Contas' not in df.columns:
            df['Contas'] = 'Nubank'
        else:
            # Preencher valores vazios/NaN com "Nubank" e remover espaços extras
            df['Contas'] = df['Contas'].fillna('Nubank')
            df['Contas'] = df['Contas'].astype(str).str.strip()
            df.loc[df['Contas'] == '', 'Contas'] = 'Nubank'
        
        return df
    except Exception as e:
        st.error(f"Erro ao carregar lançamentos: {e}")
        return pd.DataFrame()

# Função para adicionar novo lançamento
def adicionar_lancamento(data, descricao, categoria, tipo, valor, metodo, conta, status='Realizado'):
    """Adiciona um novo lançamento à planilha"""
    try:
        # Carregar o arquivo Excel
        wb = openpyxl.load_workbook(ARQUIVO_EXCEL)
        ws = wb['Lançamentos']
        
        # Encontrar a próxima linha vazia
        ultima_linha = ws.max_row + 1
        
        # Inserir dados
        ws[f'A{ultima_linha}'] = data
        ws[f'B{ultima_linha}'] = descricao
        ws[f'C{ultima_linha}'] = categoria
        ws[f'D{ultima_linha}'] = tipo
        ws[f'E{ultima_linha}'] = valor
        ws[f'F{ultima_linha}'] = metodo
        ws[f'G{ultima_linha}'] = status
        ws[f'H{ultima_linha}'] = conta
        
        # Salvar o arquivo
        wb.save(ARQUIVO_EXCEL)
        wb.close()
        
        return True
    except Exception as e:
        st.error(f"Erro ao adicionar lançamento: {e}")
        return False

# Estilo CSS personalizado
st.markdown("""
    <style>
    .main-header {
        text-align: center;
        color: #1F497D;
        margin-bottom: 30px;
    }
    .success-box {
        background-color: #d4edda;
        border: 1px solid #c3e6cb;
        color: #155724;
        padding: 12px;
        border-radius: 4px;
        margin-bottom: 20px;
    }
    .info-box {
        background-color: #d1ecf1;
        border: 1px solid #bee5eb;
        color: #0c5460;
        padding: 12px;
        border-radius: 4px;
        margin-bottom: 20px;
    }
    </style>
""", unsafe_allow_html=True)

# Título
st.markdown("<h1 class='main-header'>💰 Registro de Lançamentos Financeiros</h1>", unsafe_allow_html=True)

# Carregar configurações
config = carregar_configuracoes()

if config is None:
    st.error("Não foi possível carregar as configurações. Verifique se o arquivo Excel está correto.")
    st.stop()

# Sidebar com informações
with st.sidebar:
    st.header("📊 Informações")
    st.info("Use este aplicativo para registrar seus gastos e receitas de forma simples e rápida.")
    
    # Mostrar últimos lançamentos
    st.subheader("📝 Últimos Lançamentos")
    df_lancamentos = carregar_lancamentos()
    
    # Corrigir tipos de dados no DataFrame
    for col in df_lancamentos.select_dtypes(include=['object']).columns:
        df_lancamentos[col] = df_lancamentos[col].astype(str)
    
    if not df_lancamentos.empty:
        colunas = ['Data', 'Descrição', 'Tipo', 'Valor']
        if 'Contas' in df_lancamentos.columns:
            colunas.insert(1, 'Contas')
        ultimos = df_lancamentos.tail(5)[colunas].copy()
        st.dataframe(ultimos, use_container_width=True)
    else:
        st.write("Nenhum lançamento registrado ainda.")

# Abas principais
tab1, tab2, tab3 = st.tabs(["📥 Novo Lançamento", "📊 Visualizar Lançamentos", "📈 Resumo"])

# TAB 1: Novo Lançamento
with tab1:
    st.subheader("Adicione um novo lançamento")
    
    # Adicionar seleção de conta
    conta = st.selectbox(
        "🏦 Conta",
        options=['Nubank', 'Banco do Brasil'],
        help="Selecione a conta do lançamento"
    )
    
    col1, col2 = st.columns(2)
    
    with col1:
        data = st.date_input(
            "📅 Data do Lançamento",
            value=datetime.now(),
            help="Selecione a data do lançamento"
        )
        
        descricao = st.text_input(
            "📝 Descrição",
            placeholder="Ex: Almoço, Compras, Salário...",
            help="Descreva brevemente o lançamento"
        )
    
    with col2:
        tipo = st.radio(
            "💵 Tipo",
            options=config['tipos'],
            horizontal=True,
            help="Selecione se é receita ou despesa"
        )
    
    # Definir categorias baseadas no tipo selecionado
    if tipo == 'Receita':
        categorias_disponiveis = ['Salário', 'Extras']
    else:
        categorias_disponiveis = config['todas_categorias']
    
    with col1:
        categoria = st.selectbox(
            "🏷️ Categoria",
            options=categorias_disponiveis,
            help="Selecione a categoria do lançamento"
        )
    
    with col2:
        valor = st.number_input(
            "💰 Valor",
            min_value=0.0,
            step=0.01,
            format="%.2f",
            help="Insira o valor do lançamento"
        )
        
        # Se for Receita, método é automático (Pix/Débito)
        if tipo == 'Receita':
            st.text_input(
                "💳 Método de Pagamento",
                value="Pix/Débito",
                disabled=True,
                help="Método automático para receitas"
            )
            metodo = 'Pix/Débito'
        else:
            metodo = st.selectbox(
                "💳 Método de Pagamento",
                options=config['metodos'],
                help="Selecione o método de pagamento"
            )
    
    # Botão para adicionar
    col_btn1, col_btn2, col_btn3 = st.columns([1, 1, 2])
    
    with col_btn1:
        if st.button("✅ Adicionar Lançamento", use_container_width=True):
            if not descricao:
                st.error("⚠️ Por favor, preencha a descrição.")
            elif valor <= 0:
                st.error("⚠️ O valor deve ser maior que zero.")
            else:
                # Converter data para datetime
                data_datetime = pd.Timestamp(data)
                
                if adicionar_lancamento(data_datetime, descricao, categoria, tipo, valor, metodo, conta):
                    st.markdown(
                        f"<div class='success-box'><strong>✅ Sucesso!</strong> Lançamento adicionado com sucesso!</div>",
                        unsafe_allow_html=True
                    )
                    # Limpar cache para atualizar dados
                    st.cache_data.clear()
                else:
                    st.error("❌ Erro ao adicionar lançamento.")
    
    with col_btn2:
        if st.button("🔄 Limpar Formulário", use_container_width=True):
            st.rerun()
    
    # Resumo Financeiro Realizado vs A Transcorrer
    st.divider()
    st.subheader("💰 Resumo Financeiro")
    
    # Carregar lançamentos
    df_resumo = carregar_lancamentos()
    
    if not df_resumo.empty:
        # Data atual
        hoje = pd.Timestamp(datetime.now().date())
        
        # Calcular o último dia do mês seguinte
        if hoje.month == 12:
            proximo_mes = hoje.replace(year=hoje.year + 1, month=1, day=1)
        else:
            proximo_mes = hoje.replace(month=hoje.month + 1, day=1)
        
        # Último dia do mês seguinte
        if proximo_mes.month == 12:
            ultimo_dia_mes_seguinte = proximo_mes.replace(year=proximo_mes.year + 1, month=1, day=1) - pd.Timedelta(days=1)
        else:
            ultimo_dia_mes_seguinte = proximo_mes.replace(month=proximo_mes.month + 1, day=1) - pd.Timedelta(days=1)
        
        # Realizados: até hoje (inclusive)
        df_realizados = df_resumo[df_resumo['Data'].dt.date <= hoje.date()]
        receitas_realizadas = df_realizados[df_realizados['Tipo'] == 'Receita']['Valor'].sum()
        despesas_realizadas = df_realizados[df_realizados['Tipo'] == 'Despesa']['Valor'].sum()
        saldo_realizado = receitas_realizadas - despesas_realizadas
        
        # A transcorrer: de amanhã até o final do mês seguinte
        amanha = hoje + pd.Timedelta(days=1)
        df_transcorrer = df_resumo[
            (df_resumo['Data'].dt.date >= amanha.date()) & 
            (df_resumo['Data'].dt.date <= ultimo_dia_mes_seguinte.date())
        ]
        receitas_transcorrer = df_transcorrer[df_transcorrer['Tipo'] == 'Receita']['Valor'].sum()
        despesas_transcorrer = df_transcorrer[df_transcorrer['Tipo'] == 'Despesa']['Valor'].sum()
        
        # Exibir métricas
        col_res1, col_res2, col_res3 = st.columns(3)
        
        with col_res1:
            st.metric("✅ Saldo Realizado (até hoje)", f"R$ {saldo_realizado:,.2f}")
        
        with col_res2:
            st.metric("📅 Despesas a Transcorrer", f"R$ {despesas_transcorrer:,.2f}")
        
        with col_res3:
            st.metric("📅 Receitas a Transcorrer", f"R$ {receitas_transcorrer:,.2f}")
    else:
        st.info("📭 Nenhum lançamento registrado ainda.")

# TAB 2: Visualizar Lançamentos
with tab2:
    st.subheader("Histórico de Lançamentos")
    
    df_lancamentos = carregar_lancamentos()
    
    # Corrigir tipos de dados no DataFrame
    for col in df_lancamentos.select_dtypes(include=['object']).columns:
        df_lancamentos[col] = df_lancamentos[col].astype(str)
    
    if not df_lancamentos.empty:
        # Filtros - linha 1
        col_filt1, col_filt2, col_filt3 = st.columns(3)
        
        with col_filt1:
            tipo_filtro = st.multiselect(
                "Filtrar por Tipo",
                options=df_lancamentos['Tipo'].unique(),
                default=df_lancamentos['Tipo'].unique()
            )
        
        with col_filt2:
            categoria_filtro = st.multiselect(
                "Filtrar por Categoria",
                options=df_lancamentos['Categoria'].unique(),
                default=df_lancamentos['Categoria'].unique()
            )
        
        with col_filt3:
            # Filtro de conta (se a coluna existir)
            if 'Contas' in df_lancamentos.columns:
                conta_filtro = st.multiselect(
                    "Filtrar por Conta",
                    options=df_lancamentos['Contas'].unique(),
                    default=df_lancamentos['Contas'].unique()
                )
        
        # Filtros - linha 2
        col_filt4, col_filt5 = st.columns(2)
        
        with col_filt4:
            data_inicio = st.date_input("Data Inicial", value=df_lancamentos['Data'].min())
        
        with col_filt5:
            data_fim = st.date_input("Data Final", value=df_lancamentos['Data'].max())
        
        # Aplicar filtros
        filtro_base = (
            (df_lancamentos['Tipo'].isin(tipo_filtro)) &
            (df_lancamentos['Categoria'].isin(categoria_filtro)) &
            (df_lancamentos['Data'].dt.date >= data_inicio) &
            (df_lancamentos['Data'].dt.date <= data_fim)
        )
        
        # Adicionar filtro de conta se a coluna existir
        if 'Contas' in df_lancamentos.columns:
            filtro_base = filtro_base & (df_lancamentos['Contas'].isin(conta_filtro))
        
        df_filtrado = df_lancamentos[filtro_base]
        
        # Exibir tabela
        st.dataframe(df_filtrado, use_container_width=True, hide_index=True)
        
        # Estatísticas
        st.divider()
        st.subheader("📊 Estatísticas do Período")
        
        col_stat1, col_stat2, col_stat3 = st.columns(3)
        
        receitas = df_filtrado[df_filtrado['Tipo'] == 'Receita']['Valor'].sum()
        despesas = df_filtrado[df_filtrado['Tipo'] == 'Despesa']['Valor'].sum()
        saldo = receitas - despesas
        
        with col_stat1:
            st.metric("💚 Total Receitas", f"R$ {receitas:,.2f}")
        
        with col_stat2:
            st.metric("❤️ Total Despesas", f"R$ {despesas:,.2f}")
        
        with col_stat3:
            cor = "green" if saldo >= 0 else "red"
            st.metric("💙 Saldo", f"R$ {saldo:,.2f}", delta=None)
    else:
        st.info("📭 Nenhum lançamento registrado ainda.")

# TAB 3: Resumo
with tab3:
    st.subheader("📈 Resumo Financeiro")
    
    df_lancamentos = carregar_lancamentos()
    
    if not df_lancamentos.empty:
        # Totais gerais
        col_total1, col_total2, col_total3 = st.columns(3)
        
        total_receitas = df_lancamentos[df_lancamentos['Tipo'] == 'Receita']['Valor'].sum()
        total_despesas = df_lancamentos[df_lancamentos['Tipo'] == 'Despesa']['Valor'].sum()
        saldo_geral = total_receitas - total_despesas
        
        with col_total1:
            st.metric("💚 Total Receitas", f"R$ {total_receitas:,.2f}")
        
        with col_total2:
            st.metric("❤️ Total Despesas", f"R$ {total_despesas:,.2f}")
        
        with col_total3:
            st.metric("💙 Saldo Geral", f"R$ {saldo_geral:,.2f}")
        
        # Adicionar resumo por conta se a coluna existir
        if 'Contas' in df_lancamentos.columns:
            st.divider()
            st.subheader("🏦 Saldo por Conta")
            
            col_conta1, col_conta2 = st.columns(2)
            
            contas = df_lancamentos['Contas'].unique()
            for idx, conta in enumerate(contas):
                df_conta = df_lancamentos[df_lancamentos['Contas'] == conta]
                receitas_conta = df_conta[df_conta['Tipo'] == 'Receita']['Valor'].sum()
                despesas_conta = df_conta[df_conta['Tipo'] == 'Despesa']['Valor'].sum()
                saldo_conta = receitas_conta - despesas_conta
                
                col = col_conta1 if idx % 2 == 0 else col_conta2
                with col:
                    st.markdown(f"**{conta}**")
                    st.metric("Saldo", f"R$ {saldo_conta:,.2f}", delta=None)
                    st.caption(f"Receitas: R$ {receitas_conta:,.2f} | Despesas: R$ {despesas_conta:,.2f}")
        
        st.divider()
        
        # Gráficos
        col_chart1, col_chart2 = st.columns(2)
        
        with col_chart1:
            st.subheader("Despesas por Categoria")
            despesas_cat = df_lancamentos[df_lancamentos['Tipo'] == 'Despesa'].groupby('Categoria')['Valor'].sum().sort_values(ascending=False)
            if not despesas_cat.empty:
                st.bar_chart(despesas_cat)
            else:
                st.info("Nenhuma despesa registrada.")
        
        with col_chart2:
            st.subheader("Receitas por Categoria")
            receitas_cat = df_lancamentos[df_lancamentos['Tipo'] == 'Receita'].groupby('Categoria')['Valor'].sum().sort_values(ascending=False)
            if not receitas_cat.empty:
                st.bar_chart(receitas_cat)
            else:
                st.info("Nenhuma receita registrada.")
        
        st.divider()
        
        # Resumo por método
        st.subheader("Métodos de Pagamento Utilizados")
        metodos_uso = df_lancamentos.groupby('Método')['Valor'].sum().sort_values(ascending=False)
        if not metodos_uso.empty:
            st.bar_chart(metodos_uso)
        else:
            st.info("Nenhum método registrado.")
    else:
        st.info("📭 Nenhum lançamento registrado ainda. Comece adicionando um novo lançamento!")

# Rodapé
st.divider()
st.markdown("""
    <div style='text-align: center; color: #999; font-size: 12px; margin-top: 20px;'>
    💰 Aplicativo de Registro de Lançamentos Financeiros | Desenvolvido com Python e Streamlit
    </div>
""", unsafe_allow_html=True)
